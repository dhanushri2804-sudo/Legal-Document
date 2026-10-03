from __future__ import annotations

import os
import logging
import re
import tempfile
import unicodedata
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
from typing import Annotated

import jwt
import dotenv
from fastapi import Depends, FastAPI, File, HTTPException, Response, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.security import HTTPAuthorizationCredentials, HTTPBearer
from passlib.context import CryptContext
from pydantic import BaseModel, Field, field_validator
from sqlalchemy import func, update
from sqlalchemy.orm import Session
from docx import Document as DocxDocument
from docx.shared import Inches
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageOps, UnidentifiedImageError

from .database import Base, SessionLocal, engine, get_db
from .models import Document, DocumentVersion, Reminder, SignatureRequest, SystemAsset, User, UserLogo

dotenv.load_dotenv()
logger = logging.getLogger(__name__)

SECRET_KEY = os.getenv("JWT_SECRET_KEY", "dev-secret-key")
ALGORITHM = "HS256"
ACCESS_TOKEN_EXPIRE_MINUTES = 60 * 24

pwd_context = CryptContext(schemes=["bcrypt"], deprecated="auto")
security = HTTPBearer(auto_error=False)

app = FastAPI(title="LegalEase API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173", "http://localhost:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)


class RegisterRequest(BaseModel):
    name: str
    email: str
    password: str

    @field_validator("name")
    @classmethod
    def validate_name(cls, value):
        if not value or len(value.strip()) < 2:
            raise ValueError("Name must be at least 2 characters long.")
        return value.strip()

    @field_validator("email")
    @classmethod
    def validate_email(cls, value):
        if "@" not in value or "." not in value:
            raise ValueError("Please provide a valid email address.")
        return value.strip().lower()

    @field_validator("password")
    @classmethod
    def validate_password(cls, value):
        if len(value) < 8:
            raise ValueError("Password must be at least 8 characters long.")
        return value


class LoginRequest(BaseModel):
    email: str
    password: str


class DocumentRequest(BaseModel):
    title: str
    document_type: str
    parties: str
    terms: str
    effective_date: str
    jurisdiction: str | None = None
    additional_instructions: str | None = None
    company_name: str | None = None
    logo_filename: str | None = None

    @field_validator("title")
    @classmethod
    def validate_title(cls, value):
        if not value or len(value.strip()) < 3:
            raise ValueError("Document title is required.")
        return value.strip()

    @field_validator("document_type")
    @classmethod
    def validate_document_type(cls, value):
        if not value or len(value.strip()) < 2:
            raise ValueError("Please choose a valid document type.")
        return value.strip()

    @field_validator("parties")
    @classmethod
    def validate_parties(cls, value):
        if not value or len(value.strip()) < 3:
            raise ValueError("Parties involved is required.")
        return value.strip()

    @field_validator("terms")
    @classmethod
    def validate_terms(cls, value):
        if not value or len(value.strip()) < 20:
            raise ValueError("Terms and conditions must be at least 20 characters long.")
        return value.strip()

    @field_validator("effective_date")
    @classmethod
    def validate_effective_date(cls, value):
        if not value:
            raise ValueError("Effective date is required.")
        try:
            datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("Effective date must be a valid ISO date string.") from exc
        return value


class UpdateDocumentRequest(BaseModel):
    title: str | None = None
    document_type: str | None = None
    parties: str | None = None
    terms: str | None = None
    effective_date: str | None = None
    jurisdiction: str | None = None
    additional_instructions: str | None = None
    company_name: str | None = None
    content: str | None = None


class AssistantRequest(BaseModel):
    question: str = Field(min_length=3, max_length=2000)
    document_id: int | None = None
    document_text: str | None = Field(default=None, max_length=50000)


class SignatureRequestCreate(BaseModel):
    document_id: int
    recipient_name: str = Field(min_length=2, max_length=100)
    recipient_email: str = Field(min_length=5, max_length=255)

    @field_validator("recipient_email")
    @classmethod
    def validate_recipient_email(cls, value):
        normalized = value.strip().lower()
        if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", normalized):
            raise ValueError("A valid recipient email is required.")
        return normalized


class SignatureStatusUpdate(BaseModel):
    status: str

    @field_validator("status")
    @classmethod
    def validate_status(cls, value):
        if value not in {"pending", "signed", "declined", "cancelled"}:
            raise ValueError("Status must be pending, signed, declined, or cancelled.")
        return value


class ReminderCreate(BaseModel):
    title: str = Field(min_length=3, max_length=200)
    due_date: date
    document_id: int | None = None
    notes: str | None = Field(default=None, max_length=2000)


class ReminderUpdate(BaseModel):
    completed: bool


class UserResponse(BaseModel):
    id: int
    name: str
    email: str


class DocumentResponse(BaseModel):
    id: int
    title: str
    document_type: str
    parties: str
    terms: str
    effective_date: str
    jurisdiction: str | None
    additional_instructions: str | None
    company_name: str | None
    content: str
    created_at: str
    updated_at: str
    user_id: int


def create_access_token(data: dict):
    expires_delta = timedelta(minutes=ACCESS_TOKEN_EXPIRE_MINUTES)
    expire = datetime.now(timezone.utc) + expires_delta
    to_encode = data.copy()
    to_encode.update({"exp": expire})
    return jwt.encode(to_encode, SECRET_KEY, algorithm=ALGORITHM)


def verify_password(plain_password: str, hashed_password: str) -> bool:
    return pwd_context.verify(plain_password, hashed_password)


def get_password_hash(password: str) -> str:
    return pwd_context.hash(password)


def get_current_user(credentials: Annotated[HTTPAuthorizationCredentials | None, Depends(security)], db: Session = Depends(get_db)) -> User:
    if credentials is None or not credentials.credentials:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required.")

    token = credentials.credentials
    try:
        payload = jwt.decode(token, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = payload.get("sub")
    except jwt.PyJWTError as exc:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid or expired token.") from exc

    if user_id is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid token payload.")

    user = db.query(User).filter(User.id == int(user_id)).first()
    if user is None:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="User not found.")
    return user


def safe_document(item: Document) -> DocumentResponse:
    return DocumentResponse(
        id=item.id,
        title=item.title,
        document_type=item.document_type,
        parties=item.parties,
        terms=item.terms,
        effective_date=item.effective_date,
        jurisdiction=item.jurisdiction,
        additional_instructions=item.additional_instructions,
        company_name=item.company_name,
        content=item.content,
        created_at=item.created_at.isoformat(),
        updated_at=item.updated_at.isoformat(),
        user_id=item.user_id,
    )


def generate_fallback_document(document: DocumentRequest) -> str:
    company = document.company_name or "Your Company"
    salutations = "This draft is prepared for informational use and should be reviewed by legal counsel before execution."
    return f"{document.title}\n\nDocument Type: {document.document_type}\nEffective Date: {document.effective_date}\nJurisdiction: {document.jurisdiction or 'Not specified'}\n\nParties: {document.parties}\nCompany Name: {company}\n\nTerms and Conditions:\n{document.terms}\n\nAdditional Instructions:\n{document.additional_instructions or 'No additional instructions provided.'}\n\n{salutations}\n\nPrepared by LegalEase AI."


def generate_document_text(document: DocumentRequest) -> str:
    try:
        import google.generativeai as genai
    except ImportError:
        return generate_fallback_document(document)

    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        return generate_fallback_document(document)

    try:
        genai.configure(api_key=api_key)
        model = genai.GenerativeModel("gemini-1.5-flash")
        prompt = (
            "You are a professional legal drafting assistant. Generate a polished, clear legal document draft "
            f"based on the following inputs. Keep the tone formal and legally accurate, and include headings. "
            f"Document Type: {document.document_type}\nTitle: {document.title}\nParties: {document.parties}\n"
            f"Effective Date: {document.effective_date}\nJurisdiction: {document.jurisdiction or 'Not specified'}\n"
            f"Company Name: {document.company_name or 'Not specified'}\nTerms and Conditions: {document.terms}\n"
            f"Additional Instructions: {document.additional_instructions or 'None'}"
        )
        response = model.generate_content(prompt)
        text = getattr(response, "text", None)
        if text:
            return text
    except Exception:
        pass
    return generate_fallback_document(document)


def export_document_pdf(content: str, title: str, logo_data: bytes | None = None) -> bytes:
    if not content or not content.strip():
        raise ValueError("Document content is empty.")

    def pdf_safe_text(value: str) -> str:
        normalized = unicodedata.normalize("NFKD", value)
        return normalized.encode("latin-1", errors="replace").decode("latin-1")

    def draw_brand_mark(pdf, x: float, y: float, size: float):
        pdf.set_fill_color(11, 31, 58)
        pdf.rect(x, y, size, size, "F")
        pdf.set_draw_color(212, 175, 55)
        pdf.set_line_width(0.7)
        pdf.line(x + size * 0.22, y + size * 0.72, x + size * 0.5, y + size * 0.22)
        pdf.line(x + size * 0.5, y + size * 0.22, x + size * 0.78, y + size * 0.72)
        pdf.line(x + size * 0.29, y + size * 0.5, x + size * 0.71, y + size * 0.5)
        pdf.line(x + size * 0.5, y + size * 0.22, x + size * 0.5, y + size * 0.86)
        pdf.set_fill_color(212, 175, 55)
        pdf.ellipse(x + size * 0.2, y + size * 0.65, size * 0.14, size * 0.14, "F")
        pdf.ellipse(x + size * 0.66, y + size * 0.65, size * 0.14, size * 0.14, "F")

    def fit_image(image_width: int, image_height: int, max_width: float, max_height: float) -> tuple[float, float]:
        scale = min(max_width / image_width, max_height / image_height)
        return image_width * scale, image_height * scale

    logo_path = None
    if logo_data:
        with tempfile.NamedTemporaryFile(prefix="legalease-logo-", suffix=".png", delete=False) as logo_file:
            logo_file.write(logo_data)
            logo_path = logo_file.name

    class LegalEasePDF(FPDF):
        def header(self):
            self.set_fill_color(11, 31, 58)
            self.set_draw_color(11, 31, 58)
            if logo_path:
                with Image.open(logo_path) as logo_image:
                    logo_width, logo_height = fit_image(logo_image.width, logo_image.height, 44, 14)
                self.image(logo_path, x=20, y=10 + (14 - logo_height) / 2, w=logo_width, h=logo_height)
                text_x = 20 + logo_width + 4
            else:
                draw_brand_mark(self, 20, 9, 14)
                text_x = 38
            self.set_xy(text_x, 11)
            self.set_text_color(11, 31, 58)
            self.set_font("Helvetica", "B", 14)
            self.cell(0, 7, "LegalEase", ln=1)
            self.set_draw_color(212, 175, 55)
            self.set_line_width(0.6)
            self.line(20, 29, 190, 29)
            self.set_y(36)

        def footer(self):
            self.set_y(-12)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(100, 116, 139)
            self.cell(0, 8, f"LegalEase | Page {self.page_no()}/{{nb}}", align="C")

    pdf = LegalEasePDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(20, 20, 20)
    try:
        pdf.add_page()
        if logo_path:
            with Image.open(logo_path) as logo_image:
                body_logo_width, body_logo_height = fit_image(logo_image.width, logo_image.height, 65, 18)
            pdf.image(
                logo_path,
                x=(210 - body_logo_width) / 2,
                y=37 + (18 - body_logo_height) / 2,
                w=body_logo_width,
                h=body_logo_height,
            )
            pdf.set_y(59)
        else:
            brand_mark_size = 14
            brand_group_width = 47
            brand_group_x = (210 - brand_group_width) / 2
            draw_brand_mark(pdf, brand_group_x, 38, brand_mark_size)
            pdf.set_xy(brand_group_x + brand_mark_size + 4, 40)
            pdf.set_font("Helvetica", "B", 13)
            pdf.set_text_color(11, 31, 58)
            pdf.cell(0, 8, "LegalEase")
            pdf.set_y(59)
        pdf.set_font("Helvetica", "B", 18)
        pdf.multi_cell(0, 12, txt=pdf_safe_text(title), align="C")
        pdf.ln(6)
        pdf.set_font("Helvetica", "", 11)
        for paragraph in content.splitlines():
            clean_paragraph = pdf_safe_text(paragraph.strip())
            if not clean_paragraph:
                pdf.ln(4)
                continue
            is_heading = len(clean_paragraph) <= 100 and (
                clean_paragraph.endswith(":") or clean_paragraph.isupper()
            )
            if is_heading:
                pdf.set_font("Helvetica", "B", 12)
                pdf.multi_cell(0, 8, clean_paragraph)
                pdf.set_font("Helvetica", "", 11)
            else:
                pdf.multi_cell(0, 7, clean_paragraph)
                pdf.ln(2)
        output = pdf.output(dest="S")
        pdf_bytes = output.encode("latin-1") if isinstance(output, str) else bytes(output)
        if not pdf_bytes.startswith(b"%PDF-"):
            raise RuntimeError("PDF generator returned invalid data.")
        return pdf_bytes
    finally:
        if logo_path:
            Path(logo_path).unlink(missing_ok=True)


def export_document_docx(content: str, title: str, logo_data: bytes | None = None) -> bytes:
    doc = DocxDocument()
    header = doc.sections[0].header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = 1
    if logo_data:
        with Image.open(BytesIO(logo_data)) as logo_image:
            aspect_ratio = logo_image.width / logo_image.height
        logo_width = min(2.2, 0.7 * aspect_ratio)
        header_paragraph.add_run().add_picture(BytesIO(logo_data), width=Inches(logo_width))
    else:
        run = header_paragraph.add_run("LegalEase")
        run.bold = True
    doc.add_heading(title, level=1)
    for paragraph in content.splitlines():
        doc.add_paragraph(paragraph)
    buffer = __import__("io").BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_document_txt(content: str, title: str) -> str:
    return f"{title}\n\n{content}"


def generate_default_logo_png() -> bytes:
    scale = 3
    image = Image.new("RGBA", (192 * scale, 192 * scale), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    navy = (11, 31, 58, 255)
    gold = (212, 175, 55, 255)
    draw.rounded_rectangle((6 * scale, 6 * scale, 186 * scale, 186 * scale), radius=30 * scale, fill=navy)
    draw.line((56 * scale, 126 * scale, 96 * scale, 50 * scale), fill=gold, width=7 * scale)
    draw.line((96 * scale, 50 * scale, 136 * scale, 126 * scale), fill=gold, width=7 * scale)
    draw.line((68 * scale, 88 * scale, 124 * scale, 88 * scale), fill=gold, width=6 * scale)
    draw.line((96 * scale, 50 * scale, 96 * scale, 148 * scale), fill=gold, width=6 * scale)
    draw.ellipse((49 * scale, 116 * scale, 67 * scale, 134 * scale), fill=gold)
    draw.ellipse((125 * scale, 116 * scale, 143 * scale, 134 * scale), fill=gold)
    image = image.resize((192, 192), Image.Resampling.LANCZOS)
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def get_or_create_default_logo(db: Session) -> bytes:
    asset = db.query(SystemAsset).filter(SystemAsset.name == "legalease-default-logo").first()
    if asset:
        return asset.image_data
    asset = SystemAsset(
        name="legalease-default-logo",
        image_data=generate_default_logo_png(),
        mime_type="image/png",
        updated_at=datetime.now(timezone.utc),
    )
    db.add(asset)
    db.commit()
    db.refresh(asset)
    return asset.image_data


def find_user_document(document_id: int, user_id: int, db: Session) -> Document:
    document = db.query(Document).filter(Document.id == document_id, Document.user_id == user_id).first()
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")
    return document


def ask_gemini(prompt: str) -> str | None:
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        return None
    try:
        import google.generativeai as genai

        genai.configure(api_key=api_key)
        response = genai.GenerativeModel("gemini-1.5-flash").generate_content(prompt)
        return getattr(response, "text", None)
    except Exception:
        logger.exception("Gemini request failed.")
        return None


def analyze_contract_text(text: str) -> dict:
    paragraphs = [part.strip() for part in re.split(r"\n+|(?<=[.!?])\s+", text) if part.strip()]
    clause_terms = {
        "Payment and compensation": ("payment", "compensation", "fee", "invoice", "salary"),
        "Term and termination": ("termination", "terminate", "duration", "renewal", "expires"),
        "Confidentiality": ("confidential", "non-disclosure", "proprietary"),
        "Liability and indemnity": ("liability", "indemnif", "damages", "limitation of"),
        "Dispute resolution": ("dispute", "arbitration", "mediation", "jurisdiction", "governing law"),
        "Intellectual property": ("intellectual property", "ownership", "copyright", "work product"),
    }
    clauses = []
    lower_text = text.lower()
    for label, keywords in clause_terms.items():
        matches = [part for part in paragraphs if any(keyword in part.lower() for keyword in keywords)]
        if matches:
            clauses.append({"name": label, "excerpt": matches[0][:500]})

    missing_information = []
    if not re.search(r"\b(parties|between)\b", lower_text):
        missing_information.append("Confirm the full legal names and roles of all parties.")
    if not re.search(r"\b(effective date|commencement|start date)\b", lower_text):
        missing_information.append("Confirm the effective or commencement date.")
    if not any(term in lower_text for term in ("governing law", "jurisdiction", "laws of")):
        missing_information.append("Consider specifying governing law or jurisdiction.")
    if not any(term in lower_text for term in ("termination", "terminate", "expiration")):
        missing_information.append("Review whether the agreement needs a term and termination process.")

    date_matches = set(re.findall(
        r"\b(?:\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{4}|"
        r"(?:January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s+\d{1,2},?\s+\d{4})\b",
        text,
        re.IGNORECASE,
    ))
    extracted_dates = []
    for value in sorted(date_matches):
        parsed = None
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%B %d, %Y", "%B %d %Y"):
            try:
                parsed = datetime.strptime(value, fmt).date()
                break
            except ValueError:
                continue
        if parsed:
            extracted_dates.append({"date": parsed.isoformat(), "source_text": value})

    return {
        "summary": " ".join(paragraphs[:5])[:1200] or "No readable text was found in the document.",
        "clauses": clauses,
        "missing_information": missing_information,
        "dates": extracted_dates,
    }


@app.on_event("startup")
def startup_event():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        get_or_create_default_logo(db)
    finally:
        db.close()


@app.get("/api/health")
def health_check():
    return {"status": "ok", "message": "LegalEase backend is running"}


@app.post("/api/auth/register")
def register(request: RegisterRequest, db: Session = Depends(get_db)):
    existing = db.query(User).filter(User.email == request.email).first()
    if existing:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Email already registered.")

    user = User(name=request.name, email=request.email, password_hash=get_password_hash(request.password))
    db.add(user)
    db.commit()
    db.refresh(user)

    token = create_access_token({"sub": str(user.id)})
    return {"token": token, "user": {"id": user.id, "name": user.name, "email": user.email}}


@app.post("/api/auth/login")
def login(request: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == request.email.lower().strip()).first()
    if not user or not verify_password(request.password, user.password_hash):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid email or password.")

    token = create_access_token({"sub": str(user.id)})
    return {"token": token, "user": {"id": user.id, "name": user.name, "email": user.email}}


@app.get("/api/auth/me")
def me(current_user: User = Depends(get_current_user)):
    return {"id": current_user.id, "name": current_user.name, "email": current_user.email}


@app.get("/api/profile/logo")
def get_profile_logo(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    logo = db.query(UserLogo).filter(UserLogo.user_id == current_user.id).first()
    if logo:
        image_data = logo.image_data
        mime_type = logo.mime_type
        logo_source = "custom"
    else:
        image_data = get_or_create_default_logo(db)
        mime_type = "image/png"
        logo_source = "default"
    return Response(
        image_data,
        media_type=mime_type,
        headers={
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
            "X-Logo-Source": logo_source,
        },
    )


@app.post("/api/profile/logo")
async def save_profile_logo(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    image_bytes = await file.read(5 * 1024 * 1024 + 1)
    if len(image_bytes) > 5 * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Logo files must be 5 MB or smaller.")
    if not image_bytes:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Choose a valid PNG, JPEG, or WebP image.")

    try:
        with Image.open(BytesIO(image_bytes)) as uploaded_image:
            if uploaded_image.format not in {"PNG", "JPEG", "WEBP"}:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Choose a PNG, JPEG, or WebP image.")
            if uploaded_image.width * uploaded_image.height > 40_000_000:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Logo image dimensions must be 40 megapixels or smaller.")
            uploaded_image.verify()
        with Image.open(BytesIO(image_bytes)) as uploaded_image:
            logo_image = ImageOps.exif_transpose(uploaded_image).convert("RGBA")
            logo_image.thumbnail((1600, 800), Image.Resampling.LANCZOS)
            flattened = Image.new("RGB", logo_image.size, "white")
            flattened.paste(logo_image, mask=logo_image.getchannel("A"))
            sanitized_buffer = BytesIO()
            flattened.save(sanitized_buffer, format="PNG", optimize=True)
            sanitized_logo = sanitized_buffer.getvalue()
    except HTTPException:
        raise
    except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError) as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The logo file is invalid or could not be processed.") from exc

    logo = db.query(UserLogo).filter(UserLogo.user_id == current_user.id).first()
    if logo:
        logo.image_data = sanitized_logo
        logo.mime_type = "image/png"
        logo.updated_at = datetime.now(timezone.utc)
    else:
        db.add(UserLogo(user_id=current_user.id, image_data=sanitized_logo, mime_type="image/png"))
    db.commit()
    return {"message": "Organization logo saved.", "mime_type": "image/png"}


@app.delete("/api/profile/logo")
def delete_profile_logo(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    logo = db.query(UserLogo).filter(UserLogo.user_id == current_user.id).first()
    if logo:
        db.delete(logo)
        db.commit()
    return {"message": "Organization logo removed. PDFs will use the LegalEase logo."}


@app.get("/api/dashboard")
def dashboard(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    total_documents = db.query(func.count(Document.id)).filter(Document.user_id == current_user.id).scalar() or 0
    recent_documents = db.query(Document).filter(Document.user_id == current_user.id).order_by(Document.created_at.desc()).limit(5).all()
    stats = db.query(Document.document_type, func.count(Document.id).label("count")).filter(Document.user_id == current_user.id).group_by(Document.document_type).all()

    return {
        "welcome": f"Welcome back, {current_user.name.split()[0]}!",
        "total_documents": total_documents,
        "recent_documents": [
            {
                "id": item.id,
                "title": item.title,
                "document_type": item.document_type,
                "created_at": item.created_at.isoformat(),
            }
            for item in recent_documents
        ],
        "document_stats": [{"name": item[0], "count": item[1]} for item in stats],
    }


@app.get("/api/documents")
def get_documents(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    documents = db.query(Document).filter(Document.user_id == current_user.id).order_by(Document.created_at.desc()).all()
    return [safe_document(item) for item in documents]


@app.get("/api/documents/{document_id}")
def get_document(document_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    document = db.query(Document).filter(Document.id == document_id, Document.user_id == current_user.id).first()
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")
    return safe_document(document)


@app.post("/api/documents/generate")
def generate_document(request: DocumentRequest, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    generated_content = generate_document_text(request)
    document = Document(
        user_id=current_user.id,
        title=request.title,
        document_type=request.document_type,
        parties=request.parties,
        terms=request.terms,
        effective_date=request.effective_date,
        jurisdiction=request.jurisdiction,
        additional_instructions=request.additional_instructions,
        company_name=request.company_name,
        logo_filename=request.logo_filename,
        content=generated_content,
    )
    db.add(document)
    db.commit()
    db.refresh(document)
    return {"message": "Document generated successfully.", "document": safe_document(document)}


@app.put("/api/documents/{document_id}")
def update_document(document_id: int, request: UpdateDocumentRequest, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    document = find_user_document(document_id, current_user.id, db)

    changes = request.model_dump(exclude_unset=True, exclude_none=True)
    if any(getattr(document, field) != value for field, value in changes.items()):
        previous_versions = db.query(func.max(DocumentVersion.version_number)).filter(
            DocumentVersion.document_id == document.id
        ).scalar() or 0
        db.add(DocumentVersion(
            document_id=document.id,
            user_id=current_user.id,
            version_number=previous_versions + 1,
            title=document.title,
            content=document.content,
        ))
    for field, value in changes.items():
        if value is not None:
            setattr(document, field, value)
    document.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(document)
    return {"message": "Document updated successfully.", "document": safe_document(document)}


@app.delete("/api/documents/{document_id}")
def delete_document(document_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    document = find_user_document(document_id, current_user.id, db)
    db.query(DocumentVersion).filter(DocumentVersion.document_id == document.id).delete()
    db.query(SignatureRequest).filter(SignatureRequest.document_id == document.id).delete()
    db.execute(update(Reminder).where(Reminder.document_id == document.id).values(document_id=None))
    db.delete(document)
    db.commit()
    return {"message": "Document deleted successfully."}


@app.get("/api/documents/{document_id}/download")
def download_document(document_id: int, format: str, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    document = find_user_document(document_id, current_user.id, db)
    logo = db.query(UserLogo).filter(UserLogo.user_id == current_user.id).first()
    logo_data = logo.image_data if logo else get_or_create_default_logo(db)

    safe_title = re.sub(r"[^A-Za-z0-9._-]+", "_", document.title).strip("._-") or "document"
    common_headers = {"Content-Disposition": f'attachment; filename="{safe_title}.{format.lower()}"', "X-Content-Type-Options": "nosniff"}

    normalized_format = format.lower()
    if normalized_format == "pdf":
        if not document.content or not document.content.strip():
            logger.warning("Refusing PDF export for document %s with empty saved content.", document.id)
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document content is empty. Please save the document before downloading.")
        try:
            contents = export_document_pdf(document.content, document.title, logo_data)
        except ValueError as exc:
            logger.warning("PDF export rejected for document %s: %s", document.id, exc)
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
        except Exception as exc:
            logger.exception("PDF generation failed for document %s.", document.id)
            raise HTTPException(status_code=status.HTTP_500_INTERNAL_SERVER_ERROR, detail="Unable to generate PDF. Please try again.") from exc
        return Response(contents, media_type="application/pdf", headers={**common_headers, "Content-Disposition": f'attachment; filename="{safe_title}.pdf"'})
    if normalized_format == "docx":
        if not document.content or not document.content.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document content is empty. Please save the document before downloading.")
        contents = export_document_docx(document.content, document.title, logo_data)
        return Response(contents, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={**common_headers, "Content-Disposition": f'attachment; filename="{safe_title}.docx"'})
    if normalized_format == "txt":
        if not document.content or not document.content.strip():
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Document content is empty. Please save the document before downloading.")
        content = export_document_txt(document.content, document.title)
        return Response(content.encode("utf-8"), media_type="text/plain", headers={**common_headers, "Content-Disposition": f'attachment; filename="{safe_title}.txt"'})

    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported format. Use pdf, docx, or txt.")


@app.post("/api/documents/{document_id}/upload-logo")
def upload_logo(document_id: int, file: UploadFile = File(...), current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    document = db.query(Document).filter(Document.id == document_id, Document.user_id == current_user.id).first()
    if not document:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document not found.")

    if file and file.filename:
        document.logo_filename = file.filename
        db.commit()
        return {"message": "Logo uploaded successfully.", "filename": file.filename}
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="A valid file is required.")


@app.get("/api/documents/{document_id}/versions")
def get_document_versions(document_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    find_user_document(document_id, current_user.id, db)
    versions = db.query(DocumentVersion).filter(
        DocumentVersion.document_id == document_id,
        DocumentVersion.user_id == current_user.id,
    ).order_by(DocumentVersion.version_number.desc()).all()
    return [{
        "id": version.id,
        "version_number": version.version_number,
        "title": version.title,
        "content": version.content,
        "created_at": version.created_at.isoformat(),
    } for version in versions]


@app.post("/api/documents/{document_id}/versions/{version_id}/restore")
def restore_document_version(document_id: int, version_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    document = find_user_document(document_id, current_user.id, db)
    version = db.query(DocumentVersion).filter(
        DocumentVersion.id == version_id,
        DocumentVersion.document_id == document_id,
        DocumentVersion.user_id == current_user.id,
    ).first()
    if not version:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Document version not found.")
    latest_version = db.query(func.max(DocumentVersion.version_number)).filter(
        DocumentVersion.document_id == document_id
    ).scalar() or 0
    db.add(DocumentVersion(
        document_id=document.id,
        user_id=current_user.id,
        version_number=latest_version + 1,
        title=document.title,
        content=document.content,
    ))
    document.title = version.title
    document.content = version.content
    document.updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"message": "Document version restored.", "document": safe_document(document)}


@app.post("/api/assistant/ask")
def ask_assistant(request: AssistantRequest, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if request.document_id is not None:
        document = find_user_document(request.document_id, current_user.id, db)
        document_text = document.content
        document_title = document.title
    elif request.document_text and request.document_text.strip():
        document_text = request.document_text.strip()
        document_title = "Provided text"
    else:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Choose a document or provide text to ask about.")

    prompt = (
        "Answer the user's question using only the document excerpt. Explain legal terminology in plain language. "
        "Do not present this as legal advice; state when the document does not provide enough information.\n\n"
        f"Document ({document_title}):\n{document_text[:30000]}\n\nQuestion: {request.question}"
    )
    answer = ask_gemini(prompt)
    if answer:
        return {"answer": answer, "mode": "ai"}

    words = [word.lower() for word in re.findall(r"[A-Za-z0-9]{4,}", request.question)]
    relevant = [line.strip() for line in re.split(r"[\n.!?]+", document_text) if line.strip() and any(word in line.lower() for word in words)]
    if "summar" in request.question.lower():
        answer = "Document overview (local excerpt): " + " ".join(document_text.split()[:120])
    elif relevant:
        answer = "Relevant text from the document: " + " ".join(relevant[:3])[:2000]
    else:
        answer = "I could not find a passage that clearly answers that question in the selected text. Try asking about a specific clause or term."
    return {"answer": answer, "mode": "local", "notice": "Gemini is not configured; this is a limited text-based response, not legal advice."}


@app.post("/api/analyzer/analyze")
async def analyze_contract(file: UploadFile = File(...), current_user: User = Depends(get_current_user)):
    filename = (file.filename or "").lower()
    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Upload a PDF or DOCX file.")
    contents = await file.read(10 * 1024 * 1024 + 1)
    if len(contents) > 10 * 1024 * 1024:
        raise HTTPException(status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, detail="Files must be 10 MB or smaller.")
    try:
        if filename.endswith(".pdf"):
            from pypdf import PdfReader

            text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(contents)).pages)
        else:
            doc = DocxDocument(BytesIO(contents))
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    except Exception as exc:
        logger.info("Could not parse uploaded contract %s.", filename, exc_info=exc)
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="The uploaded file could not be read. Check that it is a valid PDF or DOCX.") from exc
    if not text.strip():
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="No readable text was found. Scanned PDFs need OCR before analysis.")

    analysis = analyze_contract_text(text[:50000])
    ai_review = ask_gemini(
        "Review this contract text for items that deserve human legal review. Return concise, clearly labeled "
        "observations only; do not claim to determine enforceability or provide legal advice.\n\n" + text[:30000]
    )
    return {
        **analysis,
        "filename": file.filename,
        "characters_analyzed": min(len(text), 50000),
        "ai_review": ai_review,
        "mode": "ai-assisted" if ai_review else "local",
        "notice": "Automated review can miss context and is not legal advice. Ask a qualified lawyer to review important agreements.",
    }


@app.get("/api/signatures")
def list_signature_requests(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    records = db.query(SignatureRequest).filter(SignatureRequest.user_id == current_user.id).order_by(SignatureRequest.created_at.desc()).all()
    return [{
        "id": item.id,
        "document_id": item.document_id,
        "document_title": item.document.title,
        "recipient_name": item.recipient_name,
        "recipient_email": item.recipient_email,
        "status": item.status,
        "created_at": item.created_at.isoformat(),
        "updated_at": item.updated_at.isoformat(),
    } for item in records]


@app.post("/api/signatures")
def create_signature_request(request: SignatureRequestCreate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    find_user_document(request.document_id, current_user.id, db)
    item = SignatureRequest(
        document_id=request.document_id,
        user_id=current_user.id,
        recipient_name=request.recipient_name.strip(),
        recipient_email=str(request.recipient_email).lower(),
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return {"id": item.id, "document_id": item.document_id, "recipient_name": item.recipient_name, "recipient_email": item.recipient_email, "status": item.status}


@app.patch("/api/signatures/{signature_id}")
def update_signature_status(signature_id: int, request: SignatureStatusUpdate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.query(SignatureRequest).filter(
        SignatureRequest.id == signature_id,
        SignatureRequest.user_id == current_user.id,
    ).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Signature request not found.")
    item.status = request.status
    item.updated_at = datetime.now(timezone.utc)
    db.commit()
    return {"id": item.id, "status": item.status, "updated_at": item.updated_at.isoformat()}


@app.get("/api/reminders")
def list_reminders(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    records = db.query(Reminder).filter(Reminder.user_id == current_user.id).order_by(Reminder.completed, Reminder.due_date).all()
    return [{
        "id": item.id,
        "document_id": item.document_id,
        "document_title": item.document.title if item.document_id else None,
        "title": item.title,
        "due_date": item.due_date.isoformat(),
        "notes": item.notes,
        "completed": item.completed,
    } for item in records]


@app.post("/api/reminders")
def create_reminder(request: ReminderCreate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    if request.document_id is not None:
        find_user_document(request.document_id, current_user.id, db)
    item = Reminder(
        user_id=current_user.id,
        document_id=request.document_id,
        title=request.title.strip(),
        due_date=request.due_date,
        notes=request.notes,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return {"id": item.id, "document_id": item.document_id, "title": item.title, "due_date": item.due_date.isoformat(), "notes": item.notes, "completed": item.completed}


@app.patch("/api/reminders/{reminder_id}")
def update_reminder(reminder_id: int, request: ReminderUpdate, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    item = db.query(Reminder).filter(Reminder.id == reminder_id, Reminder.user_id == current_user.id).first()
    if not item:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Reminder not found.")
    item.completed = request.completed
    db.commit()
    return {"id": item.id, "completed": item.completed}
